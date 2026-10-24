import csv
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from models import get_session, Expense, Category, User

class ExportManager:
    """Manage export functionality for expenses"""
    
    def __init__(self, user_id):
        self.user_id = user_id
        self.session = get_session()
    
    def export_to_csv(self, start_date=None, end_date=None, category_id=None):
        """Export expenses to CSV format"""
        query = self.session.query(
            Expense.date,
            Category.name.label('category'),
            Expense.amount,
            Expense.description,
            Expense.notes
        ).join(Category).filter(Expense.user_id == self.user_id)
        
        if start_date:
            query = query.filter(Expense.date >= start_date)
        if end_date:
            query = query.filter(Expense.date <= end_date)
        if category_id:
            query = query.filter(Expense.category_id == category_id)
        
        query = query.order_by(Expense.date.desc())
        
        expenses = query.all()
        
        # Create CSV in memory
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write header
        writer.writerow(['Date', 'Category', 'Amount', 'Description', 'Notes'])
        
        # Write data
        for expense in expenses:
            writer.writerow([
                expense.date.strftime('%Y-%m-%d'),
                expense.category,
                expense.amount,
                expense.description or '',
                expense.notes or ''
            ])
        
        return output.getvalue()
    
    def export_to_excel(self, start_date=None, end_date=None, category_id=None):
        """Export expenses to Excel format with formatting"""
        query = self.session.query(
            Expense.date,
            Category.name.label('category'),
            Category.icon.label('icon'),
            Expense.amount,
            Expense.description,
            Expense.notes
        ).join(Category).filter(Expense.user_id == self.user_id)
        
        if start_date:
            query = query.filter(Expense.date >= start_date)
        if end_date:
            query = query.filter(Expense.date <= end_date)
        if category_id:
            query = query.filter(Expense.category_id == category_id)
        
        query = query.order_by(Expense.date.desc())
        
        expenses = query.all()
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"
        
        # Header styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        # Write headers
        headers = ['Date', 'Category', 'Amount', 'Description', 'Notes']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        # Write data
        total_amount = 0
        for row, expense in enumerate(expenses, start=2):
            ws.cell(row=row, column=1, value=expense.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=f"{expense.icon} {expense.category}")
            ws.cell(row=row, column=3, value=expense.amount)
            ws.cell(row=row, column=4, value=expense.description or '')
            ws.cell(row=row, column=5, value=expense.notes or '')
            total_amount += expense.amount
        
        # Add total row
        total_row = len(expenses) + 2
        ws.cell(row=total_row, column=2, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=total_amount).font = Font(bold=True)
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 30
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    def generate_monthly_report(self, month, year):
        """Generate comprehensive monthly report"""
        from analytics import ExpenseAnalytics
        
        analytics = ExpenseAnalytics(self.user_id)
        
        # Get user info
        user = self.session.query(User).filter_by(id=self.user_id).first()
        
        # Create workbook with multiple sheets
        wb = Workbook()
        
        # Sheet 1: Summary
        ws_summary = wb.active
        ws_summary.title = "Summary"
        
        ws_summary['A1'] = f"Monthly Expense Report - {month}/{year}"
        ws_summary['A1'].font = Font(size=16, bold=True)
        
        ws_summary['A3'] = "User:"
        ws_summary['B3'] = user.username if user else "Unknown"
        
        ws_summary['A4'] = "Report Generated:"
        ws_summary['B4'] = datetime.now().strftime('%Y-%m-%d %H:%M')
        
        # Salary comparison
        comparison = analytics.get_salary_comparison(month, year)
        ws_summary['A6'] = "Monthly Salary:"
        ws_summary['B6'] = comparison['salary']
        
        ws_summary['A7'] = "Total Expenses:"
        ws_summary['B7'] = comparison['expenses']
        
        ws_summary['A8'] = "Savings:"
        ws_summary['B8'] = comparison['savings']
        ws_summary['B8'].font = Font(color="00FF00" if comparison['savings'] >= 0 else "FF0000")
        
        # Sheet 2: Category breakdown
        ws_categories = wb.create_sheet("Categories")
        ws_categories['A1'] = "Category-wise Breakdown"
        ws_categories['A1'].font = Font(size=14, bold=True)
        
        ws_categories.append(['Category', 'Amount', 'Percentage'])
        
        categories = analytics.get_category_wise_expenses()
        for cat in categories:
            ws_categories.append([
                f"{cat['icon']} {cat['name']}",
                cat['total'],
                f"{cat['percentage']}%"
            ])
        
        analytics.close()
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output
    
    def close(self):
        """Close database session"""
        self.session.close()
